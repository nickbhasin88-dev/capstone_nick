"""
Change in health aid by donor: a native (editable) PowerPoint bar chart on a true-to-scale x-axis.
Values as on the deck slide "Health aid falls about 25% by 2030" (change 2023 to 2030; source line on that slide).
Sized and placed for the right-hand chart area of that slide (about 6.2 x 3.1 in), so it pastes into place.

    python slides/build_donor_change_chart.py   -> slides/donor_change_chart.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "donor_change_chart.pptx"
DATA = [("United States", -66), ("Germany", -41), ("France", -34), ("United Kingdom", -32), ("Netherlands", -32),
        ("Japan", -15), ("Canada", 2), ("Gates Foundation", 22)]
NAVY, LIGHT, GRID, INK, MUTED = "1C3073", "8FA8E0", "D5DCEA", "1F2937", "595959"
FONT = "Arial"


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
s = prs.slides.add_slide(prs.slide_layouts[6])

cd = CategoryChartData(number_format='+0%;−0%;0%')
cd.categories = [d for d, _ in reversed(DATA)]            # bar charts draw bottom-up: reverse so the US is on top
cd.add_series("Change in health aid", [v / 100 for _, v in reversed(DATA)])
x, y, w, h = Inches(6.95), Inches(2.0), Inches(6.2), Inches(3.1)
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, cd).chart
ch.has_legend = False
ch.has_title = False
ch.font.name, ch.font.size, ch.font.color.rgb = FONT, Pt(12), rgb(INK)

plot = ch.plots[0]
plot.gap_width = 40
plot.vary_by_categories = False
ser = plot.series[0]
ser.invert_if_negative = False
for i, (_, v) in enumerate(reversed(DATA)):
    pt = ser.points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(NAVY if v < 0 else LIGHT)
    pt.format.line.fill.background()
    dpt = pt.format._element                             # the <c:dPt>; PowerPoint draws negative bars white unless
    if dpt.find(qn("c:invertIfNegative")) is None:      # each point also says invertIfNegative = 0
        inv = dpt.makeelement(qn("c:invertIfNegative"), {"val": "0"})
        dpt.find(qn("c:idx")).addnext(inv)
    lab = pt.data_label
    lab.position = XL_LABEL_POSITION.OUTSIDE_END
    lab.has_text_frame = False
    lab.show_value = True
    lab.font.name, lab.font.size, lab.font.bold, lab.font.color.rgb = FONT, Pt(13), True, rgb(INK)
    if v <= -50:                                         # longest bar: label inside its end, clear of the donor name
        lab.position = XL_LABEL_POSITION.INSIDE_END
        lab.font.color.rgb = rgb("FFFFFF")

va = ch.value_axis                                      # the x-axis (horizontal) in a bar chart
va.minimum_scale, va.maximum_scale, va.major_unit = -0.80, 0.30, 0.20
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = rgb(GRID)
va.major_gridlines.format.line.width = Pt(0.75)
va.tick_labels.number_format = '+0%;−0%;0%'
va.tick_labels.number_format_is_linked = False
va.tick_labels.font.name, va.tick_labels.font.size, va.tick_labels.font.color.rgb = FONT, Pt(10), rgb(MUTED)
va.format.line.color.rgb = rgb(GRID)
va.major_tick_mark = XL_TICK_MARK.NONE
va.has_title = True
va.axis_title.text_frame.text = "Change in health aid, 2023 to 2030"
r = va.axis_title.text_frame.paragraphs[0].runs[0]
r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(10), False, rgb(MUTED)

ca = ch.category_axis
ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW     # donor names at the far left, not at zero
ca.major_tick_mark = XL_TICK_MARK.NONE
ca.format.line.color.rgb = rgb(INK)                     # the vertical line at 0%
ca.format.line.width = Pt(1.5)
ca.tick_labels.font.name, ca.tick_labels.font.size, ca.tick_labels.font.color.rgb = FONT, Pt(14), rgb(INK)

prs.save(OUT)
print("wrote", OUT)
