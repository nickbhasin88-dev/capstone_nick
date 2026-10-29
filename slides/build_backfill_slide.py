"""
"Historically when aid has fallen, governments have not backfilled the gap": central estimates of the fiscal
response by country group with corrected 95% ranges (audit/fiscal_history/RESULTS_rerun_check.md), in the original slide's design.
Every text in Yale 2024. Estimates and ranges are read from audit/fiscal_history/rerun_check.csv (IHME per-person
columns, as the dashboard).

    python slides/build_backfill_slide.py   -> slides/backfill_slide.pptx
"""
from pathlib import Path

import pandas as pd
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from build_dashboard_segment import CB, FL, FR, FT, FW, GRID, LIGHT, MUTED, NAVY, NAVY_TEXT, Deck, line, rgb, text

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
F = "Yale 2024"
ROWS = [("All countries", "All countries (97)"), ("Upper-middle income", "Upper-middle income (32)"),
        ("Lower-middle income", "Lower-middle income (37)"), ("Low income", "Low income (26)"),
        ("Most aid-dependent third", "Most aid-dependent third (32)"),
        ("MOU countries in the panel", "MOU countries in the panel (15)")]
XMIN, XMAX, XSTEP = -0.8, 0.3, 0.2


def minus(v, d=2):
    s = f"{abs(v):.{d}f}"
    return ("−" if v < -0.0049 else "+" if v > 0.0049 else "") + s


r = pd.read_csv(ROOT / "audit" / "fiscal_history" / "rerun_check.csv")
r = r[r.data.str.startswith("IHME")].set_index("group")
for g, lab in ROWS:
    assert int(r.loc[g, "n_countries"]) == int(lab.split("(")[1].rstrip(")")), g

d = Deck(HERE / "render" / "team_capstone_deck.pptx")
d.prs.part.drop_rel(d.title_slide.rId)
d.order.remove(d.title_slide)
s = d.content("No Country Group Replaced Lost Aid With Its Own Health Spending")
for sh in s.shapes:
    if sh.name == "Title 1":
        for p in sh.text_frame.paragraphs:
            for run in p.runs:
                run.font.name, run.font.bold = F, True
    if "Slide Number" in sh.name:                         # the number is a field (a:fld), not a run
        for fld in sh._element.iter(qn("a:fld")):
            rpr = fld.find(qn("a:rPr"))
            if rpr is None:
                rpr = fld.makeelement(qn("a:rPr"), {})
                fld.insert(0, rpr)
            for tag in ("a:latin", "a:ea", "a:cs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
                el.set("typeface", F)

# subtitle
text(s, FL, FT + 90000, FW, 420000, ["Change in government health spending per US$1 of aid lost, by country group, 2002-23"],
     size=21, color=NAVY, bold=True, font=F)

# geometry
lab_w = 3500000
px0 = FL + lab_w + 120000                     # plot left
px1 = FR - 120000                             # plot right
top = FT + 760000
axis_h = 640000
key_h = 250000
bottom = CB - key_h - axis_h - 60000
rh = (bottom - top) / len(ROWS)


def X(v):
    return px0 + (v - XMIN) / (XMAX - XMIN) * (px1 - px0)


# grid and ticks
ticks = [round(XMIN + k * XSTEP, 1) for k in range(int(round((0.2 - XMIN) / XSTEP)) + 1)]      # labels to +0.2; room right for values
for t in ticks:
    if abs(t) > 1e-9:
        line(s, X(t), top, X(t), bottom, color=GRID, width=0.75)
    text(s, X(t) - 400000, bottom + 40000, 800000, 300000, [minus(t, 1) if abs(t) > 1e-9 else "0"], size=14,
         color=MUTED, font=F, align=PP_ALIGN.CENTER)
line(s, px0, bottom, px1, bottom, color=GRID, width=0.75)
text(s, px0, bottom + 330000, px1 - px0, 300000,
     ["Change in government health spending per person, per US$1 of aid lost (constant 2023 US$)"],
     size=13, color=MUTED, font=F, align=PP_ALIGN.CENTER)

# rows: light bar = 95% range, navy dot = central estimate, estimate printed at the bar's right end (as the original)
bar_h, dot, lab_gap, val_w = 230000, 240000, 90000, 900000
for k, (g, lab) in enumerate(ROWS):
    cy = top + rh * (k + 0.5)
    if k > 0:
        line(s, FL, top + rh * k, FR, top + rh * k, color=GRID, width=0.5)
    est, lo, hi = (float(r.loc[g, c]) for c in ("theta", "theta_lo", "theta_hi"))
    text(s, FL, cy - 220000, lab_w, 440000, [lab], size=18, color=NAVY_TEXT, bold=(k == 0), font=F,
         anchor=MSO_ANCHOR.MIDDLE)
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(X(lo))), Emu(int(cy - bar_h / 2)), Emu(int(X(hi) - X(lo))),
                           Emu(bar_h))
    b.fill.solid()
    b.fill.fore_color.rgb = rgb(LIGHT)
    b.line.fill.background()
    b.shadow.inherit = False
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(X(est) - dot / 2)), Emu(int(cy - dot / 2)), Emu(dot), Emu(dot))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(NAVY)
    c.line.color.rgb = rgb("FFFFFF")
    c.line.width = Pt(1.25)
    c.shadow.inherit = False
    text(s, X(max(hi, 0)) + lab_gap, cy - 220000, val_w, 440000, [minus(est)], size=18, bold=True, color=NAVY_TEXT,
         font=F, anchor=MSO_ANCHOR.MIDDLE)
line(s, X(0), top, X(0), bottom, color=NAVY, width=1.75)      # zero line on top of the bars

# key and source
text(s, FL, CB - key_h - 40000, FW, key_h,
     ["Dot = central estimate; bar = 95% range. Below 0: government health spending fell with the aid; above 0 would be "
      "replacement."],
     size=12, color=MUTED, font=F, anchor=MSO_ANCHOR.BOTTOM)
text(s, FL, CB + 40000, FW, 190000,
     ["Source: IHME Health Spending 1995-2023 (GHES per person); IHME DAH 1990-2026. Country and year fixed effects; "
      "country-clustered standard errors."],
     size=11, color="000000", font=F, anchor=MSO_ANCHOR.BOTTOM)

out = HERE / "backfill_slide.pptx"
d.save(out, ["Across 97 countries, 2002-2023, when health aid fell by a dollar per person, government health spending fell "
             "by about 19 cents over that year and the next: none of it was replaced. "
             "Low-income countries, the most aid-dependent third and the 15 MOU countries show no clear change either way. "
             "When aid rose, government spending did not move. Same pattern as Dieleman & Hanlon 2014 (Health Economics)."])
print("wrote", out)
