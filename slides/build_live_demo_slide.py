"""
One slide: the live-demo slide, without the five steps. A button that opens the dashboard, its address, a QR code
for phones, and one line on what the dashboard covers.

    python slides/build_live_demo_slide.py   -> slides/live_demo_slide.pptx   (requires segno for the QR code)
"""
from io import BytesIO
from pathlib import Path

import segno
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from build_dashboard_segment import (CB, FL, FR, FT, FW, GRID, MUTED, NAVY, NAVY_TEXT, Deck, line, rgb, text)

HERE = Path(__file__).resolve().parent
URL = "https://capstonenick-cqqwspxohi22ybxsn6ph6m.streamlit.app/"     # the link on the team deck's button

d = Deck(HERE / "render" / "team_capstone_deck.pptx")
d.prs.part.drop_rel(d.title_slide.rId)              # this file holds the one slide only
d.order.remove(d.title_slide)
s = d.content("Live demo: pick a country and a cut, see which services are at risk")

# left: the button and the address
qr_side = 2400000
lx, lw = FL, FW - qr_side - 900000
bh = 1050000
by = FT + (CB - FT) // 2 - bh // 2 - 300000
btn = s.shapes.add_shape(5, lx + 300000, by, lw - 600000, bh)        # rounded rectangle, as on the team slide
btn.adjustments[0] = 0.5
btn.fill.solid()
btn.fill.fore_color.rgb = rgb(NAVY)
btn.line.fill.background()
btn.shadow.inherit = False
tf = btn.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Open the live dashboard  ›"
r.font.name, r.font.size, r.font.bold = "Franklin Gothic Medium", Pt(32), False
r.font.color.rgb = rgb("FFFFFF")
btn.click_action.hyperlink.address = URL
text(s, lx, by + bh + 180000, lw, 340000, [URL.removeprefix("https://").rstrip("/")], size=14, color=MUTED,
     align=PP_ALIGN.CENTER)

# right: QR code for phones
buf = BytesIO()
segno.make(URL, error="m").save(buf, kind="png", scale=20, border=1, dark="#1C3073")
buf.seek(0)
qx = FR - qr_side
qy = by + bh // 2 - qr_side // 2 - 150000
s.shapes.add_picture(buf, qx, qy, qr_side, qr_side)
text(s, qx, qy + qr_side + 60000, qr_side, 320000, ["Scan to open on your phone"], size=14, color=MUTED,
     align=PP_ALIGN.CENTER)
line(s, qx - 450000, qy, qx - 450000, qy + qr_side, color=GRID, width=1.0)

# bottom: what the dashboard covers, one line
fy = CB - 420000
line(s, FL, fy - 120000, FR, fy - 120000, color=GRID, width=1.0)
text(s, FL, fy, FW, 400000,
     [[("71 countries", {"bold": True, "color": NAVY_TEXT}), ("   ·   7 Donor Scenarios   ·   4 Government Response "
                                                               "options   ·   ranges from 400 parameter draws", {})]],
     size=16, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

out = HERE / "live_demo_slide.pptx"
d.save(out, ["Click the button (or scan the code) to open the dashboard live. Suggested path: pick a country, choose a "
             "Donor Scenario, read which services and how much coverage are lost, then switch the Government Response."])
print("wrote", out)
