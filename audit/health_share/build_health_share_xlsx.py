"""
Government health spending as % of government spending, 2000-2023 (WHO GHED gghed_gge): the model's countries
(country_data/_index.csv) on the first sheet, every GHED country on the second.

    .venv/bin/python audit/health_share/build_health_share_xlsx.py  -> audit/health_share/health_share_gov_spending.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent.parent
GHED = ROOT / "datasets" / "WHO_GHED" / "GHED_data_2026-10-14.xlsx"
INDEX = ROOT / "country_data" / "_index.csv"
OUT = Path(__file__).resolve().parent / "health_share_gov_spending.xlsx"
MOU15 = {"BWA", "BDI", "CMR", "SWZ", "ETH", "KEN", "LSO", "LBR", "MDG", "MWI", "MOZ", "NGA", "RWA", "SLE", "UGA"}
ABUJA = 0.15
MAIN = "Model Countries"

F = "Arial"
HEAD = PatternFill("solid", fgColor="1C3073")
REPEAT = PatternFill("solid", fgColor="EDEDED")
INPUT = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="D5DCEA"))

g = pd.read_excel(GHED, sheet_name="Data", usecols=lambda c: c in ("location", "code", "year", "gghed_gge"))
names = g.drop_duplicates("code").set_index("code")["location"]
w = g.pivot(index="code", columns="year", values="gghed_gge") / 100
w = w[w.notna().any(axis=1)]
years = list(w.columns)
model = pd.read_csv(INDEX)
missing = sorted(set(model.recipient_isocode) - set(w.index))

wb = Workbook()


def sheet(ws, title, codes, abuja_ref, abuja_cell=False):
    ws["A1"] = title
    ws["A1"].font = Font(name=F, bold=True, size=12)
    ws["A2"] = "Abuja target"
    ws["A2"].font = Font(name=F, bold=True)
    if abuja_cell:
        ws["B2"] = ABUJA
        ws["B2"].font = Font(name=F, color="0000FF")
        ws["B2"].fill = INPUT
        ws["B2"].comment = Comment("Abuja Declaration (2001): 15% of government spending on health. "
                                   "Edit to test another target; both sheets use this cell.", "model")
    else:
        ws["B2"] = f"={abuja_ref}"
        ws["B2"].font = Font(name=F, color="008000")
    ws["B2"].number_format = "0.0%"
    ws["C2"] = ("Grey cell = same value as the year before (likely carried forward by WHO, not a new figure). "
                "Blank = no WHO estimate. See Notes.")
    ws["C2"].font = Font(name=F, italic=True, size=9, color="595959")
    hr = 4
    heads = ["Country", "ISO3", "MOU 15"] + [str(y) for y in years] + ["Latest", "Latest Year", "Avg 2019-2023",
                                                                       "Room to Abuja target"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(hr, j, h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    y0 = 4
    first_y, last_y = get_column_letter(y0), get_column_letter(y0 + len(years) - 1)
    c19 = get_column_letter(y0 + years.index(2019))
    lat, laty, avg, room = (get_column_letter(y0 + len(years) + k) for k in range(4))
    for i, iso in enumerate(sorted(codes, key=lambda c: names[c])):
        r = hr + 1 + i
        ws.cell(r, 1, names[iso]).font = Font(name=F)
        ws.cell(r, 2, iso).font = Font(name=F)
        ws.cell(r, 3, "yes" if iso in MOU15 else "").font = Font(name=F)
        prev = None
        for j, y in enumerate(years):
            v = w.loc[iso, y]
            c = ws.cell(r, y0 + j, None if pd.isna(v) else float(v))
            c.number_format = "0.0%"
            c.font = Font(name=F, color="0000FF")
            if prev is not None and not pd.isna(v) and round(v, 6) == round(prev, 6):
                c.fill = REPEAT
            prev = v
        ws[f"{lat}{r}"] = f"=LOOKUP(9.99E+307,{first_y}{r}:{last_y}{r})"
        ws[f"{laty}{r}"] = f"=LOOKUP(9.99E+307,{first_y}{r}:{last_y}{r},{first_y}${hr}:{last_y}${hr})"
        ws[f"{avg}{r}"] = f"=IFERROR(AVERAGE({c19}{r}:{last_y}{r}),\"\")"
        ws[f"{room}{r}"] = f"=MAX(0,$B$2-{lat}{r})"
        for col in (lat, avg, room):
            ws[f"{col}{r}"].number_format = "0.0%"
        for col in (lat, laty, avg, room):
            ws[f"{col}{r}"].font = Font(name=F)
        ws[f"{laty}{r}"].alignment = Alignment(horizontal="center")
        for j in range(1, len(heads) + 1):
            ws.cell(r, j).border = THIN
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 7
    for j in range(y0, y0 + len(years)):
        ws.column_dimensions[get_column_letter(j)].width = 7
    for col in (lat, laty, avg, room):
        ws.column_dimensions[col].width = 12
    ws.row_dimensions[hr].height = 30
    ws.freeze_panes = ws.cell(hr + 1, y0)
    ws.auto_filter.ref = f"A{hr}:{room}{hr + len(codes)}"


m = wb.active
m.title = MAIN
sheet(m, f"Government health spending as % of government spending, model countries ({len(model) - len(missing)}), "
         "2000-2023 (WHO GHED, GGHE-D % GGE)", [c for c in model.recipient_isocode if c in w.index], None, abuja_cell=True)
sheet(wb.create_sheet("All GHED Countries"), f"Government health spending as % of government spending, all WHO GHED "
      f"countries ({len(w)}), 2000-2023", list(w.index), f"'{MAIN}'!$B$2")

# ---------------------------------------------------------------- spending by category (IMF COFOG, cofog_all/)
COFOG = ROOT / "cofog_all"
CAT = "Spending by Category"
FUNCS = [("GF01", "General Public Services (excl. Interest)"), ("GF0170", "Interest on Public Debt"), ("GF02", "Defence"),
         ("GF03", "Public Order and Safety"), ("GF04", "Economic Affairs"), ("GF05", "Environmental Protection"),
         ("GF06", "Housing and Community Amenities"), ("HEALTH", "Health"), ("GF08", "Recreation, Culture and Religion"),
         ("GF09", "Education"), ("GF10", "Social Protection"), ("UNCL", "Unclassified")]
HEALTH = [("GF0712", "Medical Products"), ("GF072", "Outpatient Services"), ("GF073", "Hospital Services"),
          ("GF0740", "Public Health Services"), ("GF075", "Health R&D"), ("GF0760", "Other Health"),
          ("GF07X", "Health (No Breakdown)")]
cf = pd.concat(pd.read_csv(f) for f in sorted(COFOG.glob("*.csv")))
cf["code"] = cf.code.replace({"GF07": "GF07X", "GF07_rest": "GF07X"})
cv = cf.pivot_table(index=["iso3", "year"], columns="code", values="pct_outlays", aggfunc="sum") / 100
cov = cf.groupby(["iso3", "year"]).coverage.first()
cf_missing = sorted(set(model.recipient_isocode) - set(cf.iso3))

c = wb.create_sheet(CAT)
c["A1"] = (f"Government spending by category, % of total outlays (IMF Government Finance Statistics, COFOG), "
           f"{cf.iso3.nunique()} countries, {cf.year.min()}-{cf.year.max()}")
c["A1"].font = Font(name=F, bold=True, size=12)
c["A2"] = ("One row per country-year with data. Health = sum of its sub-items (formula). Total should be 100%. "
           "WHO Health Share = the GHED figure for the same year from the 'All GHED Countries' sheet, for comparison.")
c["A2"].font = Font(name=F, italic=True, size=9, color="595959")
hr = 4
lead = ["Country", "ISO3", "MOU 15", "Year", "Coverage"]
heads = lead + [lab for _, lab in FUNCS] + ["Total"] + [lab for _, lab in HEALTH] + ["WHO Health Share (GHED)"]
col = {h: get_column_letter(j) for j, h in enumerate(heads, 1)}
for j, h in enumerate(heads, 1):
    x = c.cell(hr, j, h)
    x.font = Font(name=F, bold=True, color="FFFFFF")
    x.fill = HEAD
    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
keys = sorted(cv.index, key=lambda k: (names.get(k[0], k[0]), k[1]))
g_rng = f"'All GHED Countries'!$D$5:${get_column_letter(3 + len(years))}${4 + len(w)}"
g_iso = f"'All GHED Countries'!$B$5:$B${4 + len(w)}"
g_hdr = f"'All GHED Countries'!$D$4:${get_column_letter(3 + len(years))}$4"
for i, (iso, yr) in enumerate(keys):
    r = hr + 1 + i
    vals = [names.get(iso, iso), iso, "yes" if iso in MOU15 else "", str(yr), cov[(iso, yr)]]
    for j, v in enumerate(vals, 1):
        c.cell(r, j, v).font = Font(name=F)
    for code, lab in FUNCS + HEALTH:
        if code == "HEALTH":
            continue
        v = cv.loc[(iso, yr)].get(code, float("nan"))
        x = c[f"{col[lab]}{r}"]
        x.value = None if pd.isna(v) else float(v)
        x.font = Font(name=F, color="0000FF")
        x.number_format = "0.0%"
    h0, h1 = col[HEALTH[0][1]], col[HEALTH[-1][1]]
    c[f"{col['Health']}{r}"] = f"=SUM({h0}{r}:{h1}{r})"
    c[f"{col['Total']}{r}"] = f"=SUM({col[FUNCS[0][1]]}{r}:{col[FUNCS[-1][1]]}{r})"
    c[f"{col['WHO Health Share (GHED)']}{r}"] = (f"=IFERROR(IF(INDEX({g_rng},MATCH($B{r},{g_iso},0),MATCH($D{r},{g_hdr},0))=\"\",\"\",INDEX({g_rng},MATCH($B{r},{g_iso},0),MATCH($D{r},{g_hdr},0))),\"\")")
    for k in ("Health", "Total", "WHO Health Share (GHED)"):
        c[f"{col[k]}{r}"].number_format = "0.0%"
        c[f"{col[k]}{r}"].font = Font(name=F, bold=(k == "Health"), color="008000" if k.startswith("WHO") else "000000")
    for j in range(1, len(heads) + 1):
        c.cell(r, j).border = THIN
last_row = hr + len(keys)
c.column_dimensions["A"].width = 22
c.column_dimensions["E"].width = 16
for h in heads[5:]:
    c.column_dimensions[col[h]].width = 11
c.column_dimensions["B"].width = c.column_dimensions["C"].width = c.column_dimensions["D"].width = 7
c.row_dimensions[hr].height = 58
c.freeze_panes = "F5"
c.auto_filter.ref = f"A{hr}:{col[heads[-1]]}{last_row}"

# latest year per country, looked up from the category sheet with formulas
L = wb.create_sheet("Latest Year by Category")
L["A1"] = "Government spending by category, latest year with COFOG data for each country (% of total outlays)"
L["A1"].font = Font(name=F, bold=True, size=12)
L["A2"] = "Each value is looked up from 'Spending by Category' for the country's latest year (formulas)."
L["A2"].font = Font(name=F, italic=True, size=9, color="595959")
lh = ["Country", "ISO3", "MOU 15", "Latest Year", "Coverage"] + [lab for _, lab in FUNCS] + ["Total",
                                                                                          "WHO Health Share (GHED)"]
for j, h in enumerate(lh, 1):
    x = L.cell(hr, j, h)
    x.font = Font(name=F, bold=True, color="FFFFFF")
    x.fill = HEAD
    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
rng = lambda h: f"'{CAT}'!${col[h]}$5:${col[h]}${last_row}"          # noqa: E731
isos = sorted(cv.index.get_level_values(0).unique(), key=lambda k: names.get(k, k))
for i, iso in enumerate(isos):
    r = hr + 1 + i
    for j, v in enumerate([names.get(iso, iso), iso, "yes" if iso in MOU15 else ""], 1):
        L.cell(r, j, v).font = Font(name=F)
    L[f"D{r}"] = f"=TEXT(SUMPRODUCT(MAX(({rng('ISO3')}=$B{r})*({rng('Year')}+0))),\"0\")"
    key = f"MATCH(1,INDEX(({rng('ISO3')}=$B{r})*({rng('Year')}=$D{r}),0),0)"
    L[f"E{r}"] = f"=INDEX({rng('Coverage')},{key})"
    for j, h in enumerate(lh[5:], 6):
        x = L.cell(r, j, f"=IFERROR(IF(INDEX({rng(h)},{key})=\"\",\"\",INDEX({rng(h)},{key})),\"\")")
        x.number_format = "0.0%"
        x.font = Font(name=F, color="008000", bold=(h == "Health"))
    for j in (4, 5):
        L.cell(r, j).font = Font(name=F, color="008000")
    for j in range(1, len(lh) + 1):
        L.cell(r, j).border = THIN
L.column_dimensions["A"].width = 22
L.column_dimensions["E"].width = 16
for j in range(6, len(lh) + 1):
    L.column_dimensions[get_column_letter(j)].width = 11
L.column_dimensions["B"].width = L.column_dimensions["C"].width = 7
L.column_dimensions["D"].width = 8
L.row_dimensions[hr].height = 58
L.freeze_panes = "F5"
L.auto_filter.ref = f"A{hr}:{get_column_letter(len(lh))}{hr + len(isos)}"

n = wb.create_sheet("Notes")
notes = [
    ("Indicator", "Domestic General Government Health Expenditure (GGHE-D) as % of General Government Expenditure (GGE); "
                  "WHO GHED variable gghed_gge. World Bank WDI SH.XPD.GHED.GE.ZS republishes the same series."),
    ("Source", "WHO Global Health Expenditure Database, December 2025 update (data to 2023), downloaded 14 Oct 2026: "
               "https://apps.who.int/nha/database (file datasets/WHO_GHED/GHED_data_2026-10-14.xlsx)."),
    ("Countries", f"'{MAIN}': the dashboard's {len(model)} countries (country_data/_index.csv)"
                  + (f"; not in GHED: {', '.join(missing)}" if missing else "")
                  + ". 'All GHED Countries': every country with at least one value. 'MOU 15' marks the 15 countries "
                    "in the 2027 MOU reallocation."),
    ("Domestic only", "Government spending from its own revenue. Donor money, including aid passed through the "
                      "government budget, is not in the numerator."),
    ("Grey cells", "Same value as the year before. Runs of identical values (e.g. 2022 = 2023 in most countries) are "
                   "probably WHO carrying the last estimate forward, not new data."),
    ("Latest / Latest Year", "Last non-blank year in the row and its year (formulas)."),
    ("Room to Abuja target", f"Abuja target (cell B2 on '{MAIN}', used by both sheets) minus Latest, floored at 0."),
    ("Differences from the dashboard", "The dashboard's fiscal panel shows the WDI copy of this series (latest year); its "
                                       "path panel uses an IMF-based share, which can differ (audit item I-14)."),
    ("Spending by Category: source", "IMF Government Finance Statistics, expenditure by function of government (COFOG), "
                                     "as stored for the dashboard's spending-mix chart in cofog_all/<ISO3>.csv (no raw "
                                     "download is on disk; see build_sections.py). Shares are % of total outlays."),
    ("Spending by Category: coverage", f"{cf.iso3.nunique()} countries; years vary by country (some have one year). "
                                       f"Model countries with no COFOG data: {', '.join(cf_missing)}. The level of "
                                       "government varies: 'Central government (budgetary)' leaves out states, local "
                                       "governments and social security funds, so it understates categories run "
                                       "locally (often health and education) and is not comparable with "
                                       "'General government' rows."),
    ("Interest", "Interest on public debt is shown on its own and is not also inside General Public Services (as on the "
                 "dashboard)."),
    ("Health vs WHO share", "COFOG health is spending by function from the government budget, and can include donor "
                            "money that passes through the budget; the WHO (GHED) share counts government money from "
                            "domestic sources only, for general government. The two can differ a lot; the last column "
                            "puts them side by side."),
    ("Total", "Sum of the categories; should be 100%. A few country-years add to slightly more (up to 104.8%) in the "
              "IMF data; the dashboard rescales those to 100%, this sheet does not."),
    ("Built by", "audit/health_share/build_health_share_xlsx.py"),
]
n["A1"] = "Notes"
n["A1"].font = Font(name=F, bold=True, size=12)
for i, (k, v) in enumerate(notes, 3):
    n.cell(i, 1, k).font = Font(name=F, bold=True)
    c = n.cell(i, 2, v)
    c.font = Font(name=F)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    n.cell(i, 1).alignment = Alignment(vertical="top")
n.column_dimensions["A"].width = 28
n.column_dimensions["B"].width = 110

wb.save(OUT)
print("wrote", OUT, "| model countries", len(model), "missing from GHED:", missing, "| GHED countries", len(w))
